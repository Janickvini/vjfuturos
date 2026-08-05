import openpyxl
import json
import datetime
import os
import sys

# Caminhos dos arquivos
base_dir = os.path.dirname(os.path.abspath(__file__))

# 1. Procura primeiro por base.xlsx na pasta do dashboard
excel_dir = base_dir
excel_files = [f for f in os.listdir(excel_dir) if f.endswith('.xlsx') and not f.startswith('~$') and f == 'base.xlsx']

# 2. Se não achar, procura qualquer .xlsx na pasta do dashboard
if not excel_files:
    excel_files = [f for f in os.listdir(excel_dir) if f.endswith('.xlsx') and not f.startswith('~$')]

# 3. Se ainda não achar, procura na pasta antiga BasePPDef
if not excel_files:
    excel_dir = os.path.abspath(os.path.join(base_dir, "..", "BasePPDef"))
    if os.path.exists(excel_dir):
        excel_files = [f for f in os.listdir(excel_dir) if f.endswith('.xlsx') and not f.startswith('~$')]

if not excel_files:
    print("Erro: Nenhum arquivo .xlsx encontrado no diretório do dashboard ou em BasePPDef.")
    sys.exit(1)

# Ordena e seleciona o arquivo Excel alvo
excel_files.sort()
excel_file_name = excel_files[-1] 
wb_path = os.path.join(excel_dir, excel_file_name)

print(f"Lendo dados de: {wb_path}...")

try:
    wb = openpyxl.load_workbook(wb_path, data_only=True)
except Exception as e:
    print(f"Erro ao abrir a planilha: {e}")
    sys.exit(1)

def clean_row(row, headers):
    item = {}
    for idx, cell_val in enumerate(row):
        if idx < len(headers):
            header = headers[idx]
            if header is None:
                continue
            
            # Formatação de datas
            if isinstance(cell_val, datetime.datetime) or isinstance(cell_val, datetime.date):
                cell_val = cell_val.strftime("%Y-%m-%d")
            
            # Limpeza de strings
            if isinstance(cell_val, str):
                cell_val = cell_val.strip()
                
            item[header] = cell_val
    return item

# 1. Processar aba Stakeholders
if "Stakeholders" not in wb.sheetnames:
    print("Erro: Aba 'Stakeholders' não encontrada no arquivo Excel.")
    sys.exit(1)

sh_sheet = wb["Stakeholders"]
sh_rows = list(sh_sheet.iter_rows(values_only=True))
sh_headers = sh_rows[0]
cleaned_sh_headers = []
for h in sh_headers:
    if h is not None:
        h_str = str(h).strip()
        if h_str.startswith("="):
            cleaned_sh_headers.append(None)
        else:
            cleaned_sh_headers.append(h_str)
    else:
        cleaned_sh_headers.append(None)

stakeholders = []
for row in sh_rows[1:]:
    if any(cell is not None for cell in row):
        item = clean_row(row, cleaned_sh_headers)
        if item.get("Instituição"):
            stakeholders.append(item)

# 2. Processar aba Relatórios
if "Relatórios" not in wb.sheetnames:
    print("Erro: Aba 'Relatórios' não encontrada no arquivo Excel.")
    sys.exit(1)

rep_sheet = wb["Relatórios"]
rep_rows = list(rep_sheet.iter_rows(values_only=True))
rep_headers = rep_rows[0]
cleaned_rep_headers = []
for h in rep_headers:
    if h is not None:
        h_str = str(h).strip()
        if h_str.startswith("="):
            cleaned_rep_headers.append(None)
        else:
            cleaned_rep_headers.append(h_str)
    else:
        cleaned_rep_headers.append(None)

reports = []
for row in rep_rows[1:]:
    if any(cell is not None for cell in row):
        item = clean_row(row, cleaned_rep_headers)
        if item.get("Nome do Estudo"):
            reports.append(item)

data = {
    "stakeholders": stakeholders,
    "reports": reports
}

# Salva o arquivo data.js atualizado
out_path = os.path.join(base_dir, "data.js")
with open(out_path, "w", encoding="utf-8") as f:
    f.write("// Dados exportados automaticamente de " + excel_file_name + "\n")
    f.write("const baseData = ")
    json.dump(data, f, ensure_ascii=False, indent=2)
    f.write(";\n")

print("\n--- Atualização Concluída com Sucesso! ---")
print(f"Arquivo lido: {excel_file_name}")
print(f"Stakeholders carregados: {len(stakeholders)}")
print(f"Relatórios carregados: {len(reports)}")
print(f"Dados salvos em: {out_path}")
